"""
报告下载文件：每次运行把当天的结果导出成 CSV / Excel / PDF，放在 docs/downloads/<日期>/ 下面，
GitHub Pages 会直接提供下载；报告页面上的"📥 下载报告"区块列出最近 KEEP_REPORT_DAYS 天的文件。

- 同一天跑多次，后面的会覆盖前面的 → 每天保留的是当天最后一次运行的结果
- 超过 KEEP_REPORT_DAYS 个交易日的文件夹会被自动删掉
- 另外生成一份"近 N 天合并"的 Excel (每天一个工作表) 和 CSV (多一列日期)
"""
import csv
import json
import math
import os
import re
import shutil

DOWNLOADS_DIR = os.path.join("docs", "downloads")
KEEP_REPORT_DAYS = 7  # 保留最近几个"报告日" (只有交易日才有报告，所以约等于一周半的日历天数)
COMBINED_BASENAME = "bursa-report-last7"

COLUMNS = ["类型", "代码", "名称", "价格", "涨跌%", "成交量", "相对量", "RSI", "SAR", "EMA20", "50日均线", "信号", "AI点评"]
DATE_DIR_RE = re.compile(r"^\d{4}-\d{2}-\d{2}$")


# ---------- 数据整理 ----------

def _num(v, ndigits):
    """None / NaN 一律变成 None (导出成空白)。NaN 写进 xlsx 会让 Excel 报文件损坏，写进 CSV 会变成 "nan"。"""
    if v is None:
        return None
    try:
        v = float(v)
    except (TypeError, ValueError):
        return None
    return None if math.isnan(v) or math.isinf(v) else round(v, ndigits)


def change_pct(close, prev_close):
    """跟网页表格同一套算法：昨收先四舍五入到 3 位，避免价格没变却算出 -0.00%。"""
    prev = round(prev_close, 3) if prev_close else None
    return round((close - prev) / prev * 100, 2) if prev else 0.0


def build_rows(stocks):
    """把 main() 里的 stocks 整理成导出用的行：信号在前，其余按成交量从高到低 (跟网页顺序一致)。"""
    signals, others = [], []
    for s in stocks:
        d = s.get("data")
        if not d or d.get("low_volume"):
            continue
        row = {
            "类型": "信号" if s["matched"] else "表格",
            "代码": s["symbol"].split(".")[0],
            "名称": s["name"],
            "价格": _num(d["close"], 3),
            "涨跌%": _num(change_pct(d["close"], d["prev_close"]), 2),
            "成交量": int(d["volume"]),
            "相对量": _num(d.get("rel_volume"), 2),
            "RSI": _num(d.get("rsi"), 2),   # 价格 14 天都没动的股票 RSI 是 NaN (0/0)
            "SAR": "多头" if d["sar_bullish_now"] else "空头",
            "EMA20": _num(d.get("ema20_latest"), 3),
            "50日均线": _num(d.get("sma50"), 3),
            "信号": s.get("reason") or "",
            "AI点评": s.get("ai_comment") or "",
        }
        (signals if s["matched"] else others).append(row)
    others.sort(key=lambda r: r["成交量"], reverse=True)
    return signals + others


def report_date_of(stocks, fallback):
    """报告日期取行情数据里最新一根日线的日期 (FORCE_RUN 周末跑也不会生成一个"周六"的文件)。"""
    dates = [s["data"]["candles"][-1]["time"] for s in stocks
             if s.get("data") and not s["data"].get("low_volume") and s["data"].get("candles")]
    return max(dates) if dates else fallback


# ---------- CSV ----------

def write_csv(path, rows, with_date=False):
    cols = (["日期"] if with_date else []) + COLUMNS
    # utf-8-sig: 带 BOM，Excel 直接双击打开中文不会乱码
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=cols, extrasaction="ignore")
        w.writeheader()
        for r in rows:
            w.writerow({k: ("" if r.get(k) is None else r.get(k)) for k in cols})


# ---------- Excel ----------

def _fill_sheet(ws, rows):
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    ws.append(COLUMNS)
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="2F3E4E")
    for cell in ws[1]:
        cell.font, cell.fill = header_font, header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    up, down, muted = Font(color="0CA30C"), Font(color="D03B3B"), Font(color="898781")
    signal_fill = PatternFill("solid", fgColor="EEEAF8")
    for r in rows:
        ws.append([r.get(c) for c in COLUMNS])
        row_idx = ws.max_row
        # 涨跌% 在 Excel 里存成小数 + 百分比格式，这样还能直接排序/计算
        chg_cell = ws.cell(row=row_idx, column=COLUMNS.index("涨跌%") + 1)
        if isinstance(chg_cell.value, (int, float)):
            chg = chg_cell.value
            chg_cell.value = chg / 100
            chg_cell.number_format = "+0.00%;-0.00%;0.00%"
            chg_cell.font = up if chg > 0 else down if chg < 0 else muted
        ws.cell(row=row_idx, column=COLUMNS.index("价格") + 1).number_format = "0.000"
        ws.cell(row=row_idx, column=COLUMNS.index("EMA20") + 1).number_format = "0.000"
        ws.cell(row=row_idx, column=COLUMNS.index("50日均线") + 1).number_format = "0.000"
        ws.cell(row=row_idx, column=COLUMNS.index("成交量") + 1).number_format = "#,##0"
        ws.cell(row=row_idx, column=COLUMNS.index("相对量") + 1).number_format = "0.00"
        sar_cell = ws.cell(row=row_idx, column=COLUMNS.index("SAR") + 1)
        sar_cell.font = up if sar_cell.value == "多头" else down
        if r["类型"] == "信号":
            for c in range(1, len(COLUMNS) + 1):
                ws.cell(row=row_idx, column=c).fill = signal_fill
        ws.cell(row=row_idx, column=COLUMNS.index("AI点评") + 1).alignment = Alignment(wrap_text=True, vertical="top")

    widths = {"类型": 7, "代码": 8, "名称": 16, "价格": 9, "涨跌%": 9, "成交量": 14, "相对量": 8,
              "RSI": 7, "SAR": 7, "EMA20": 9, "50日均线": 10, "信号": 30, "AI点评": 50}
    for i, c in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(c, 10)
    ws.freeze_panes = "A2"  # 往下滚时表头固定
    if ws.max_row > 1:
        ws.auto_filter.ref = ws.dimensions  # 表头带筛选按钮


def write_xlsx(path, days):
    """days: [(日期, rows), ...]，每一天一个工作表 (单日文件就只有一个)。"""
    from openpyxl import Workbook
    wb = Workbook()
    wb.remove(wb.active)
    for date, rows in days:
        _fill_sheet(wb.create_sheet(title=date), rows)
    wb.save(path)


# ---------- PDF ----------

_EMOJI_RE = re.compile(r"[\U00010000-\U0010FFFF\u2600-\u27BF\uFE0F]")


def _plain(text):
    """PDF 用的内置中文字体没有 emoji，去掉 🎯🚨 这类字符，避免显示成方块。"""
    return _EMOJI_RE.sub("", text or "").strip()


def _fmt(v, pattern):
    return "—" if v is None else pattern.format(v)


def write_pdf(path, date, rows, generated_at):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    from xml.sax.saxutils import escape

    # reportlab 自带的中文 CID 字体，不用往仓库里放 10MB 的字体文件
    font = "STSong-Light"
    if font not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(UnicodeCIDFont(font))

    title_style = ParagraphStyle("t", fontName=font, fontSize=16, leading=20, spaceAfter=2)
    sub_style = ParagraphStyle("s", fontName=font, fontSize=9, leading=12, textColor=colors.HexColor("#52514e"))
    h_style = ParagraphStyle("h", fontName=font, fontSize=12, leading=16, spaceBefore=8, spaceAfter=4)
    body_style = ParagraphStyle("b", fontName=font, fontSize=9, leading=13)
    cell_style = ParagraphStyle("c", fontName=font, fontSize=8, leading=9)

    signals = [r for r in rows if r["类型"] == "信号"]
    others = [r for r in rows if r["类型"] != "信号"]
    up, down, muted = colors.HexColor("#0ca30c"), colors.HexColor("#d03b3b"), colors.HexColor("#898781")

    story = [
        Paragraph(f"马股自动分析报告 {date}", title_style),
        Paragraph(f"生成时间: {generated_at} (MYT)　|　信号 {len(signals)} 支　|　其余股票 {len(others)} 支", sub_style),
        Spacer(1, 4 * mm),
        Paragraph(f"信号 ({len(signals)})", h_style),
    ]
    if not signals:
        story.append(Paragraph("今日无符合条件的股票。", body_style))
    for r in signals:
        story.append(Paragraph(
            f"<b>{escape(r['名称'])}</b> {r['代码']}　现价 {_fmt(r['价格'], '{:.3f}')}　"
            f"RSI(14) {_fmt(r['RSI'], '{:.2f}')}　50日均线 {_fmt(r['50日均线'], '{:.3f}')}　{escape(_plain(r['信号']))}", body_style))
        if r["AI点评"]:
            story.append(Paragraph(f"AI 点评: {escape(_plain(r['AI点评']))}", sub_style))
        story.append(Spacer(1, 2 * mm))

    story.append(Paragraph(f"其余股票 ({len(others)}，按成交量从高到低)", h_style))
    header = ["#", "代码", "名称", "价格", "涨跌%", "成交量", "相对量", "RSI", "SAR", "EMA20"]
    data = [header]
    styles = []
    for i, r in enumerate(others, start=1):
        chg = r["涨跌%"] or 0.0
        data.append([
            str(i), r["代码"], Paragraph(escape(r["名称"]), cell_style), _fmt(r["价格"], "{:.3f}"),
            f"{chg:+.2f}%" if chg else "0.00%", f"{r['成交量']:,}",
            _fmt(r["相对量"], "{:.2f}"), _fmt(r["RSI"], "{:.1f}"), r["SAR"], _fmt(r["EMA20"], "{:.3f}"),
        ])
        styles.append(("TEXTCOLOR", (4, i), (4, i), up if chg > 0 else down if chg < 0 else muted))
        styles.append(("TEXTCOLOR", (8, i), (8, i), up if r["SAR"] == "多头" else down))

    table = Table(data, repeatRows=1, colWidths=[10 * mm, 16 * mm, 48 * mm, 20 * mm, 20 * mm,
                                                  30 * mm, 18 * mm, 16 * mm, 16 * mm, 20 * mm])
    table.setStyle(TableStyle([
        ("FONTNAME", (0, 0), (-1, -1), font),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2f3e4e")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f4f4f1")]),
        ("ALIGN", (0, 0), (0, -1), "RIGHT"),
        ("ALIGN", (3, 0), (7, -1), "RIGHT"),
        ("ALIGN", (9, 0), (9, -1), "RIGHT"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LINEBELOW", (0, 0), (-1, -1), 0.25, colors.HexColor("#e1e0d9")),
        ("TOPPADDING", (0, 0), (-1, -1), 1),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 1),
    ] + styles))
    story.append(table)

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont(font, 7)
        canvas.setFillColor(colors.HexColor("#898781"))
        canvas.drawString(12 * mm, 7 * mm, "本报告仅供技术分析参考，不构成投资建议。版权所有 CJA231，保留一切权利，未经授权禁止转载。")
        canvas.drawRightString(doc.pagesize[0] - 12 * mm, 7 * mm, f"第 {doc.page} 页")
        canvas.restoreState()

    doc = SimpleDocTemplate(path, pagesize=landscape(A4), leftMargin=12 * mm, rightMargin=12 * mm,
                            topMargin=12 * mm, bottomMargin=14 * mm,
                            title=f"马股自动分析报告 {date}", author="Bursa Bot")
    doc.build(story, onFirstPage=footer, onLaterPages=footer)


# ---------- 主入口 ----------

def _files_for(date):
    base = f"bursa-report-{date}"
    return {"csv": f"{date}/{base}.csv", "xlsx": f"{date}/{base}.xlsx", "pdf": f"{date}/{base}.pdf"}


def export_downloads(stocks, today, generated_at, downloads_dir=DOWNLOADS_DIR):
    """
    导出当天的 CSV/Excel/PDF，删掉超出保留天数的旧文件夹，再生成"近 N 天合并"文件。
    返回给网页用的清单: {"days": [{"date", "generated_at", "files", "count", "signals"}...(新→旧)],
                        "combined": {"xlsx", "csv"} 或 None}
    """
    date = report_date_of(stocks, today)
    rows = build_rows(stocks)

    day_dir = os.path.join(downloads_dir, date)
    os.makedirs(day_dir, exist_ok=True)
    files = _files_for(date)
    # data.json 是合并文件的数据来源 (不用再去解析 CSV/Excel)
    with open(os.path.join(day_dir, "data.json"), "w", encoding="utf-8") as f:
        json.dump({"date": date, "generated_at": generated_at, "rows": rows}, f, ensure_ascii=False)
    write_csv(os.path.join(downloads_dir, files["csv"]), rows)
    write_xlsx(os.path.join(downloads_dir, files["xlsx"]), [(date, rows)])
    write_pdf(os.path.join(downloads_dir, files["pdf"]), date, rows, generated_at)

    # 只处理名字是日期的文件夹，其他东西不碰
    day_dirs = sorted((d for d in os.listdir(downloads_dir)
                       if DATE_DIR_RE.match(d) and os.path.isdir(os.path.join(downloads_dir, d))), reverse=True)
    for old in day_dirs[KEEP_REPORT_DAYS:]:
        shutil.rmtree(os.path.join(downloads_dir, old))
    kept = day_dirs[:KEEP_REPORT_DAYS]

    days, combined_days = [], []
    for d in kept:
        data_path = os.path.join(downloads_dir, d, "data.json")
        if not os.path.exists(data_path):
            continue
        with open(data_path, encoding="utf-8") as f:
            payload = json.load(f)
        combined_days.append((d, payload["rows"]))
        days.append({
            "date": d,
            "generated_at": payload.get("generated_at", ""),
            "files": _files_for(d),
            # 给网页上的日期选择条显示"这一天有什么": 共几支、信号是哪几支
            "count": len(payload["rows"]),
            "signals": [r["名称"] for r in payload["rows"] if r["类型"] == "信号"],
        })

    combined = None
    if combined_days:
        combined = {"xlsx": f"{COMBINED_BASENAME}.xlsx", "csv": f"{COMBINED_BASENAME}.csv"}
        write_xlsx(os.path.join(downloads_dir, combined["xlsx"]), combined_days)
        write_csv(os.path.join(downloads_dir, combined["csv"]),
                  [dict(r, 日期=d) for d, rs in combined_days for r in rs], with_date=True)

    return {"days": days, "combined": combined, "keep_days": KEEP_REPORT_DAYS}
